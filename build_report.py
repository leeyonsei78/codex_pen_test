from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from xml.sax.saxutils import escape

ROOT=Path(__file__).parent
OUT=ROOT/'deliverables'
OUT.mkdir(exist_ok=True)
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.7)
sec.left_margin=Cm(1.9); sec.right_margin=Cm(1.9)
sec.footer_distance=Cm(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 s=doc.styles[name]; s.font.name='맑은 고딕'; s.font.color.rgb=RGBColor(0,0,0)
 s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'맑은 고딕')
 s.paragraph_format.space_after=Pt(7)
 s.paragraph_format.line_spacing=1.15
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Title'].font.size=Pt(23)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(4)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.core_properties.title='라이나생명 홈페이지 보안 점검 통합 보고서'
doc.core_properties.subject='금융권 침해사고 비교와 보안 점검 결과 및 후속 계획'
doc.core_properties.author='보안 점검 보고서'

def p(t='',bold=False,size=None):
 x=doc.add_paragraph(); r=x.add_run(t); r.bold=bold
 if size:r.font.size=Pt(size)
 return x
def h(t,l=1):doc.add_heading(t,l)
def page():doc.add_page_break()
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=v
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=str(v)
 for ri,row in enumerate(t.rows):
  trPr=row._tr.get_or_add_trPr()
  trPr.append(OxmlElement('w:cantSplit'))
  if ri==0:trPr.append(OxmlElement('w:tblHeader'))
  for i,c in enumerate(row.cells):
   c.width=Cm(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'DDE7EF' if ri==0 else ('F4F7F9' if ri%2==0 else 'FFFFFF'));pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders)
   mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.line_spacing=1.1
    for r in para.runs:r.font.size=Pt(9);r.bold=ri==0
 p('',size=2)
 return t

def diagram(title,nodes,edges,height=220):
 p(title,True)
 ns='xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"'
 xml=f'<w:pict {ns}><v:group style="width:480pt;height:{height}pt" coordorigin="0,0" coordsize="960,{height*2}">'
 for i,(x,y,w,hh,label,fill) in enumerate(nodes):
  paras=''.join('<w:p><w:pPr><w:jc w:val="center"/><w:spacing w:after="0"/></w:pPr><w:r><w:rPr><w:rFonts w:ascii="Malgun Gothic" w:hAnsi="Malgun Gothic" w:eastAsia="맑은 고딕"/><w:sz w:val="19"/></w:rPr><w:t>'+escape(line)+'</w:t></w:r></w:p>' for line in label.split('\n'))
  xml+=f'<v:roundrect id="diag{len(doc.paragraphs)}node{i}" style="position:absolute;left:{x};top:{y};width:{w};height:{hh}" arcsize="0.08" fillcolor="{fill}" strokecolor="#8B9BA8"><v:textbox inset="5pt,5pt,5pt,5pt"><w:txbxContent>{paras}</w:txbxContent></v:textbox></v:roundrect>'
 for i,(x1,y1,x2,y2,dashed) in enumerate(edges):
  xml+=f'<v:line style="position:absolute" from="{x1},{y1}" to="{x2},{y2}" strokecolor="#596A77" strokeweight="1.2pt"><v:stroke endarrow="block" dashstyle="{"dash" if dashed else "solid"}"/></v:line>'
 xml+='</v:group></w:pict>'
 para=doc.add_paragraph();para.paragraph_format.keep_with_next=True;para.add_run()._r.append(parse_xml(xml))
 p('편집 방법  Word에서 도형을 선택해 문구와 크기 및 위치를 수정할 수 있습니다. 이미지로 삽입하지 않았습니다.',size=8)

footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('라이나생명 보안 점검 통합 보고서  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)

doc.add_paragraph('라이나생명 홈페이지\n보안 점검 통합 보고서','Title')
p('금융권 침해사고 비교와 점검 결과 및 후속 조치 계획',size=12)
p('작성 기준  2026년 10월 29일 KST    점검 증적  2026년 10월 28일\n수신 대상  정보보호 책임자와 홈페이지 개발 및 운영 담당자',size=9)
p('최근 금융권 침해사고에서 드러난 외부 업무 접점과 고객정보 보호 문제를 기준으로 라이나생명 공개 홈페이지 점검 결과를 검토한다. 공개 자료에서 확인된 사실, 조건부 위험, 추가 검증이 필요한 통제를 구분하고 개선의 우선순위와 완료 기준을 제시한다.')
h('1 최근 금융권 침해사고')
h('은행권 사고와 공통 점검 방향',2)
p('2026년 10월 14일 금융위원회는 신한·국민·부산·하나은행 등을 포함한 16개 금융회사의 침해사고를 확인했다고 발표했다. 당시 집계된 정보유출은 약 10만 건이며, 동일 공격 주체나 AI 활용 여부와 개별 침입경로는 조사 중이었다. 이 숫자는 10월 29일의 최종 피해 규모가 아니다. [S1]')
table(['사례와 기준 시점','확인 또는 보도된 내용','라이나에 적용할 확인 질문'],[
 ['신한은행\n9월 30일 인지 관련 10월 2일 보도','대출모집인용 조회 시스템에서 약 2만 5천 건의 고객정보 유출 보도. 고객번호를 이용한 다른 서비스 조회 정황도 보도됨. [S2]','모집인·제휴 채널과 API 사이에 요청별 인증과 객체 소유권 검사가 있는가'],
 ['KB국민·하나·부산은행\n10월 14일 당국 발표','침해사고 확인 대상에 포함. 개별 시스템의 최종 원인은 해당 조사 결과로 구분해야 함. [S1]','대표 홈페이지 이외의 직원·협력업체 접점이 자산 목록과 점검 범위에 포함되는가']
],[3.4,6.5,7.3])
p('은행 사례의 교훈은 한 화면의 취약점 확인에 그치지 않는 것이다. 외부 접점의 인증, 후속 조회 API의 권한, 대량 조회 제한과 탐지를 함께 확인해야 한다. 이는 사례를 바탕으로 한 점검 권고이며 라이나에서 동일한 결함을 확인했다는 뜻은 아니다.')

page();h('미래에셋생명과 증권업권 사고',1)
h('미래에셋생명',2)
p('10월 25일 이상 징후를 인지하고 26일 새벽 일부 고객정보 유출을 확인했다는 회사 설명이 보도됐다. 홈페이지와 고객 청약시스템 일부가 침해 대상으로 파악됐고, 10월 28일 보도 시점에는 홈페이지 접근 제한과 복구·추가 침입 점검이 이어졌다. 구체적인 유출 인원과 정보 항목은 조사 중이었다. [S3]')
p('보험 청약정보에는 민감한 자료가 포함될 수 있으나, 이번 공개 자료만으로 병력·질병정보가 실제 유출됐다고 단정할 수 없다. 라이나는 청약·첨부서류·상담·고객 조회 연계의 데이터 흐름과 접근 권한을 별도로 검증해야 한다.')
h('대신증권',2)
p('대신증권은 10월 17일 투자일임 비대면 계좌개설 연계서버의 비정상 접근을 발견했고, 10월 28일 총 38,490명의 개인·신용정보 유출을 발표한 것으로 보도됐다. 고객명·이메일·주소·통신 ID·은행명과 암호화된 계좌번호 및 휴대전화번호 등이 포함됐다. 접근 통제와 외부 연계 시스템 점검을 강화하겠다고 밝혔다. [S4]')
p('일부 항목이 암호화됐다는 사실로 유출 전체의 피해가 없다고 판단할 수 없다. 핵심 거래시스템과 분리된 연계서버에서도 고객정보가 유출될 수 있다는 점이 라이나의 제휴·청약 API와 서비스 계정 권한 검토에 직접적인 비교 기준이 된다.')
h('미래에셋자산운용과 회사 구분',2)
p('10월 28일 미래에셋자산운용의 일부 웹 시스템에서 비정상 접근 정황과 차단 조치가 보도됐다. 당시 회사는 개인신용정보 등의 유출이 없는 것으로 파악된다고 설명했다. 이 사건은 미래에셋생명 사고 및 대신증권 사고와 구분한다. 본 비교에서 미래에셋증권의 고객정보 유출을 확인된 사건으로 기재하지 않는다. [S5]')
table(['비교 대상','실제 점검 범위를 넓혀야 하는 영역','확정하지 않은 사항'],[
 ['은행','모집인 업무와 후속 고객 조회 API','최종 침입 원인과 동일 공격자 여부'],
 ['미래에셋생명','홈페이지에서 청약·파일로 이어지는 권한과 데이터 흐름','최종 피해 규모와 병력 등 개별 유출 항목'],
 ['대신증권','외부 계좌개설 연계서버와 고객 데이터 접근 범위','라이나와 동일 제품·취약점 사용 여부']
],[3,8.3,5.9])
p('사고 요약은 발표·보도 시점의 상태다. 회사 설명을 인용한 언론 보도는 독립적인 포렌식 결론과 구분했다. AI 자동화나 특정 CVE를 라이나의 직접 원인으로 대입하지 않았다.',size=9)

page();h('2 종합 판단과 점검 범위')
p('공개 홈페이지의 브라우저 보안과 클라이언트 코드에서 개선이 필요한 근거를 확보했다. 특히 API 오류의 HTML 출력 경로와 인증 헤더 목적지 검사 공백을 우선 수정 대상으로 본다. 다만 실제 원격 XSS·토큰 유출·서버 침입·고객정보 유출은 입증하지 못했다.',True)
p('관측한 HTTPS 연결은 정상이나, 서버 인증·객체 권한·DB 접근·패치·탐지 운영의 증빙이 없다. 따라서 라이나 전체의 안전성, 유사 대규모 유출의 불가능성 또는 출시 가능 여부를 승인할 수 없다. 침해사고 발생 여부를 판단한 포렌식 보고서도 아니다.')
diagram('도식 1 점검 범위와 남아 있는 확인 영역',[
 (15,35,280,115,'공개 홈페이지와 참조 JS\nHTML 및 응답 헤더\n수집 증적 검토 완료','#E4EDF4'),
 (345,35,280,115,'브라우저 데이터 흐름\n선택 구문 로컬 재현\n실제 브라우저는 미검증','#E4EDF4'),
 (675,35,270,115,'인증 및 업무 API\n서버 소스와 전용 계정 필요\n실제 통제 미확인','#F3EEE3'),
 (345,235,600,100,'청약과 DB 및 내부 업무 시스템\n자산 목록과 설정 및 로그 필요  |  미확인','#F3EEE3')
],[(295,90,345,90,False),(625,90,675,90,True),(805,150,805,235,True)],190)
p('실선은 자료 검토 또는 로컬 재현 범위의 연결이며, 점선은 별도 증빙과 검증이 필요한 경계다. 시스템의 실제 네트워크 구성을 확정한 그림은 아니다.',size=9)
table(['수행한 작업','수행하지 않은 작업'],[
 ['정상 GET으로 공개 HTML과 직접 참조 JS 수집\n저장된 응답·헤더·해시·번들 일부 분석\n합성 데이터로 선택 원본 구문 재현\n수정 참고 코드 회귀 테스트','로그인과 실고객 정보 접근\nAPI 공격·서버 침투·부하 시험\n브라우저에서 실제 XSS 또는 헤더 전송 검증\n원본 앱 수정·운영 배포·서버 설정 검사']
],[8.6,8.6])
p('점검 산출물 기준  저장소 커밋 a7d99d3d751e484267f25ffd5cfd03e48526c650. 이번 통합 작성 과정에서 라이나 운영 서비스에 대한 추가 점검 요청은 수행하지 않았다. [E1~E5]',size=9)

page();h('3 상세 점검 결과')
p('F 번호는 기존 보고서와 일치한다. 등급은 조치 우선순위용 잠정 평가이며 CVSS나 침입 성공 판정이 아니다. F-07은 F-03의 출력 경로를 보강한 내용으로 중복 취약점 수로 합산하지 않는다.')
h('F 01 구버전 jQuery 이중 참조',2)
p('확인  홈과 로그인 HTML에서 3.7.1 뒤에 1.12.4 참조가 있고 실제 응답 파일의 버전 주석도 일치했다. 1.12.4는 공식 CVE-2020-11022 및 CVE-2020-11023 영향 범위에 포함된다. [E1][S6]')
p('판단  잠정 중 위험. 비신뢰 HTML이 DOM 조작 함수에 전달되는 실제 호출 경로와 최종 전역 버전은 미확인이다. 구버전 존재만으로 XSS 성공 또는 서버 침투를 확정하지 않는다.')
p('조치와 완료 기준  중복 참조를 제거하고 보안 지원 버전으로 통합한다. 잠금 파일과 SBOM을 확보하고 브라우저 네트워크·실행 버전 및 인증·전자서명·팝업 플러그인 회귀 결과를 남긴다.')
h('F 02 프레임 제한과 CSP 미관측',2)
p('확인  수집한 홈·로그인·404 응답에서 CSP 및 X-Frame-Options를 관측하지 못했다. 실제 클릭재킹은 재현하지 않았다. [E1]')
p('판단  프레임 제한은 잠정 중 위험의 개선 항목이다. CSP 부재는 추가 방어층의 공백이며 그 자체가 XSS 입증은 아니다.')
p('조치와 완료 기준  업무상 허용 부모 origin을 정하고 최종 CDN·프록시 응답에 frame-ancestors를 적용한다. 전체 스크립트 CSP는 Report-Only로 호환성을 확인한 후 강화한다. 정상·오류 응답과 제휴 iframe 동작을 검증한다.')
h('F 03 자체 HTML 방어 함수 오류',2)
p('확인  블랙리스트에 없는 태그·이벤트 문자열 통과, 따옴표 엔티티 오타, 이중 인코딩을 로컬 4건으로 재현했다. 함수의 실제 호출부 연결은 확인되지 않았다. [E1]')
p('판단과 조치  문맥에 맞지 않는 방어를 신뢰하지 않도록 일반 메시지는 텍스트로 출력한다. 리치 HTML만 별도 정화한다. 함수 결함과 F-07의 HTML 출력 경로를 구분하여 원본 소스에서 연결 여부를 추적한다.')
h('F 04 내부 서비스 주소 공개',2)
p('확인  공개 런타임 설정에 내부 서비스 이름·주소·포트가 포함됐다. 잠정 낮음이며 내부망 접근이나 SSRF 성립은 미입증이다. [E1]')
p('조치와 완료 기준  내부 전용 설정을 서버 영역으로 분리한다. 배포 산출물과 공개 초기 상태에서 불필요한 내부 설정이 제거됐는지 재확인한다. 공개 주소만으로 내부 라우팅과 권한을 추정하지 않는다.')

page();h('인증정보와 HTML 출력 관련 결과')
h('F 05 JavaScript에서 접근 가능한 토큰 저장',2)
p('확인  공개 번들에서 idToken과 refreshToken을 sessionStorage 및 JavaScript 쿠키로 복사하는 코드를 확인했다. 합성 값으로 저장 구문을 재현했다. 실제 로그인 토큰과 브라우저 쿠키 수용 결과는 관측하지 않았다. [E2]')
p('판단  조건부 중 위험의 설계 검토 사항이다. XSS 또는 신뢰한 외부 스크립트가 침해되면 토큰 접근 피해를 키울 수 있다. 서버의 수명·회전·폐기와 토큰 사용 정책은 미확인이다.')
p('조치와 완료 기준  구조가 허용되면 BFF 또는 서버 세션과 서버 발급 HttpOnly 쿠키를 검토한다. JavaScript 쿠키 문자열에 HttpOnly를 추가하는 방식으로 해결할 수 없다. SPA를 유지하면 짧은 토큰 수명·회전·재사용 탐지·로그아웃 폐기 및 인증 화면 외부 스크립트 최소화를 검증한다. 쿠키 인증 전환 시 CSRF 정책도 함께 설계한다. [S8]')
h('F 06 캠페인 URL의 과도한 저장과 로그',2)
p('확인  utm_source 부분 문자열 검사와 전체 쿼리 저장으로 비UTM 값과 fragment가 localStorage 및 콘솔 인자에 포함되는 동작을 합성 URL 2건으로 재현했다. [E2]')
p('판단  잠정 낮음이며 실제 URL에 민감정보가 포함되는지에 따라 재평가한다. 실제 개인정보가 저장됐거나 외부 분석 서비스로 전송됐다는 증거는 없다.')
p('조치와 완료 기준  승인된 캠페인 필드만 정확한 이름으로 파싱하고 길이·문자 제한을 적용한다. raw query와 디버그 로그를 제거하고 기존 저장 데이터의 정리 정책을 정한다. URL 자체에 인증정보와 개인정보를 넣지 않도록 한다. 참고 코드 테스트 8개가 통과했으나 운영 반영은 하지 않았다.')
h('F 07 API 오류 메시지의 HTML 출력 경로',2)
p('확인  API 오류 message가 공통 인터셉터에서 Vuex 상태를 거쳐 전역 메시지의 innerHTML용 데이터로 전달된다. 선택한 원본 구문을 합성 응답으로 실행했을 때 마크업 문자열이 유지됐다. [E3]')
p('판단  수정 우선순위 높음. HTML 전달 경로는 확인됐으나 공격자 입력의 서버 반영, 브라우저 스크립트 실행 및 토큰 악용의 전체 경로는 미입증이다. F-03과 관련된 보강 근거다.')
p('조치와 완료 기준  일반 메시지를 텍스트로 출력하고 승인 errorCode를 고정 안내 문구에 매핑한다. 원본 message와 stackTrace는 UI에 전달하지 않는다. 서버 입력에서 메시지까지의 흐름과 실제 DOM을 격리 테스트에서 확인한다. 참고 코드 테스트 6개가 통과했으나 앱 통합·배포는 하지 않았다.')

page();h('토큰 회전과 요청 목적지 관련 결과')
h('F 08 새 refresh token을 이전 값으로 덮어쓰기',2)
p('확인  특정 401 갱신 분기에서 새 응답 refreshToken에 이전 값을 다시 대입했다. 합성 새 값과 이전 값을 사용한 로컬 재현에서도 이전 값이 선택됐다. [E3]')
p('판단  서버 계약 확인 전 취약점으로 확정하지 않는다. 서버가 회전하면 새 토큰 폐기와 재사용 탐지 충돌을 유발할 수 있고, 기존 값 재사용을 허용하는 계약이면 의도된 동작일 수 있다.')
p('조치와 완료 기준  인증 서버의 회전 계약을 확보한다. 새 토큰 채택, 이전 값 재사용 거부, 동시 갱신과 다중 탭, 로그아웃·만료 후 거부를 전용 계정으로 검증한다. 코드 대입 결과만으로 서버 재사용을 허용한다고 판단하지 않는다.')
h('F 09 Axios의 교차 출처 XSRF 헤더 처리',2)
p('확인  번들 메타데이터의 Axios 0.21.4와 withCredentials 조건에서 교차 출처에도 쿠키 값을 XSRF 헤더로 복사하는 구문을 확인했다. 원본 선택 구문을 로컬 모형으로 재현했다. [E3][S7]')
p('판단  의존성 개선 필요. 실제 XSRF 쿠키, 호출 옵션, 공격자 제어 목적지 및 브라우저 요청 조건은 미확인이다. 서버 측 Axios 버전이나 SSRF·RCE를 추정할 수 없다.')
p('조치와 완료 기준  현재 보안 권고를 반영한 지원 버전으로 교체하고 실제 의존성 트리와 잠금 파일을 점검한다. 허용 API에만 필요한 인증·XSRF 헤더가 전달되는지 검증한다. 과거 개별 CVE 수정 버전을 현재의 권장 버전으로 단정하지 않는다.')
h('F 10 인증 헤더 부착 전 origin 검사 공백',2)
p('확인  요청 URL에 /api/ 또는 /public/ 구간이 있으면 정확한 origin 비교 없이 Authorization을 붙이는 구문이 있다. Axios는 외부 절대 URL을 baseURL로 제한하지 않았다. 합성 URL 5건의 메모리 요청 설정에서 외부 origin에도 합성 토큰이 붙는 경우를 재현했다. [E4]')
p('판단  조건부 인증정보 노출 위험이며 수정 우선순위 높음. 공격자의 URL 인자 제어 지점, 실제 헤더 전송 및 CORS·preflight는 미검증이다. 실제 유출이나 계정 탈취 성공으로 해석하지 않는다.')
p('조치와 완료 기준  헤더를 붙이기 전에 최종 URL의 scheme·host·port가 승인 HTTPS origin과 정확히 일치하는지 확인한다. 검사한 URL을 실제 전송에도 사용한다. 참고 검사 함수 테스트 6개가 통과했으며 원본 앱 통합은 미수행이다. F-09의 라이브러리 XSRF 처리와 별도 로직이다.')

page();h('4 금융권 사고 대비 안전성 평가')
p('비교는 통제 영역별 증빙의 수준을 평가한다. 다른 회사의 사고가 라이나에서 그대로 재현된다는 의미가 아니며, 미확인은 통제 부재나 통과 어느 쪽으로도 처리하지 않는다.')
table(['비교 영역','라이나에서 확보한 근거','안전성 판단과 추가 확인'],[
 ['외부 업무 접점\n은행·증권 공통','www와 참조 JS만 점검','전체 안전성 판단 보류. 청약·모집인·제휴·모바일·관리·배포 접점을 목록화하고 데이터 및 내부망 연결 확인'],
 ['패치와 EOS\n금융위 공통 점검','jQuery 1.12.4와 Axios 0.21.4 관측','프런트 개선 필요. 서버·WAS·OS·보안장비 패치와 지원 종료 여부는 자산별 증빙 필요'],
 ['인증과 객체 권한\n은행 후속 조회','토큰 저장과 갱신 분기 확인','서버 통제 미확인. 매 요청 인증·역할·객체·필드 권한과 인증 누락 경로 검증'],
 ['사용자 세션 피해','F-05·F-07·F-10 조건부 경로','우선 수정 필요. 실제 입력 도달·브라우저 동작·목적지 제어와 토큰 재사용 방어 확인'],
 ['청약과 고객 데이터\n미래에셋생명','서버·DB·파일 설정 없음','민감정보 및 대량 유출 방어 판단 보류. 파일 권한·DB 최소 권한·암호화·키 분리·조회 한도 확인'],
 ['외부 연계서버\n대신증권','연계 API의 서버 통제 미점검','인증정보 범위와 연계 계정 권한 확인. 한 서버 침해가 타 업무 데이터로 확산되는지 검증'],
 ['탐지와 사고 대응','CDN 흔적만 관측','24시간 관제와 실제 차단 능력 미확인. 합성 이벤트로 경보·담당자 인지·회수까지 확인'],
 ['HTTPS와 전송 보호','관측 연결 TLS 1.3 및 인증서 검증 정상','해당 연결의 긍정적 근거. 전 지점·구형 프로토콜·갱신 체계·접근 권한 안전성은 별도 확인']
],[3.4,5.0,8.8])
h('종합 안전성 판단',2)
p('브라우저 측 개선 필요성이 확인됐고 핵심 서버 통제는 확인 공백이 크다. 안전성 점수나 사고 확률을 산출할 근거가 부족하다. 라이나의 사고 발생 또는 대규모 유출 가능성이 입증된 것은 아니지만, 현재 결과만으로 유사 사고가 불가능하다고 결론낼 수도 없다.')
p('정상 관측도 보존한다. 로그아웃 코드에 서버 메서드 호출이 있고 일부 인증 오류는 토큰 상태를 비운다. production/staging에서 오류 세부 UI 노출을 제한하는 분기도 있다. 이들은 서버의 실제 토큰 폐기나 API 응답 최소화를 보증하지 않는다. [E2~E4]')

page();h('5 조건부 위험 경로와 차단 지점')
diagram('도식 2 오류 메시지에서 계정 피해까지의 확인 경계',[
 (20,10,280,95,'외부 입력 또는 비신뢰 데이터\nAPI message 반영 여부 미확인','#F3EEE3'),
 (355,10,280,95,'API 오류 message\n인터셉터에서 선택','#E4EDF4'),
 (680,10,260,95,'Vuex 저장\nHTML 전달 확인','#E4EDF4'),
 (680,185,260,95,'전역 메시지 innerHTML\n브라우저 실행 미검증','#E4EDF4'),
 (355,185,280,95,'페이지 JavaScript 실행\nXSS 성립 조건 미확인','#F3EEE3'),
 (20,185,280,95,'JS 접근 가능 토큰\n저장 코드 확인 F 05','#E4EDF4'),
 (20,350,615,75,'서버 토큰 악용과 고객정보 접근  |  실제 성공 미입증','#F3EEE3')
],[(300,55,355,55,True),(635,55,680,55,False),(810,105,810,185,False),(680,230,635,230,True),(355,230,300,230,True),(160,280,160,350,True)],225)
p('실선은 정적 흐름 또는 선택 구문 로컬 재현을 뜻한다. 점선은 아직 충족 여부를 확인하지 않은 조건이다. innerHTML까지의 연결을 확인했다고 해서 스크립트 실행과 계정 피해까지 입증한 것은 아니다. [E3]')
table(['차단 위치','우선 적용할 통제','검증할 결과'],[
 ['메시지 출력 전','텍스트 렌더링과 승인 코드 매핑','비실행 마커가 HTML 요소로 해석되지 않으며 정상 업무 안내 유지'],
 ['네트워크 요청 전','정확한 승인 origin 검사 F-10','외부 origin 입력은 전송 전 거부되고 인증 헤더 미부착'],
 ['토큰 저장과 갱신','서버 세션 또는 제한된 수명·회전·폐기','이전·만료·회수 토큰 거부와 정상 갱신 성공'],
 ['API와 데이터 접근','객체·역할·필드 권한 및 누적 조회 제한','다른 합성 고객 데이터 차단과 경보·회수 확인']
],[3.3,6.5,7.4])
p('F-10은 이 도식과 별도의 요청 목적지 경로다. URL 입력 제어가 가능한 호출부가 있는지 추적하고, 실제 브라우저 및 모의 API로 헤더 전달 여부를 검증해야 한다. 여러 조건부 경로를 하나의 성공한 공격으로 이어 붙여 평가하지 않는다.')

page();h('6 추가 확인 항목과 합격 기준')
p('격리된 테스트 환경과 전용 A·B 계정 및 합성 데이터를 먼저 준비한다. API 명세를 기준으로 주체·기능·객체·메서드·필드별 허용 정답을 정하고 정상 허용 요청과 금지 요청을 함께 검사한다. 아래 실제 서버 검증 상태는 모두 미수행이다. [E5]')
table(['검증 영역','실행할 확인','합격 기준과 보관할 증빙'],[
 ['인증과 관리자','미인증·만료·잘못된 토큰, 일반 계정의 관리 기능 요청','민감 데이터와 상태 변경 없음. 정책에 맞는 거부 응답 및 서버 인증 로그'],
 ['객체와 필드 권한','A가 B 소유 합성 객체를 조회·변경하고 제한 필드 요청','서버에서 객체·행위·필드별 거부. 숨김 UI가 아닌 API 결과와 DB 변경 없음 확인'],
 ['세션 수명주기','로그아웃·회수·만료·이전 refresh 재사용·동시 갱신','승인한 폐기 기한 내 거부, 정상 갱신 성공. 인증 서버 계약·시간 기록·감사 로그'],
 ['조회와 반출 제한','단건 반복·페이지 이동·동시 작업·내보내기 누적량','계정·업무·데이터량 한도와 경보 기한을 사전 합의하고 차단·알림 확인'],
 ['서버 입력과 파일','스키마·길이·타입·업로드 저장·다운로드 권한·파라미터 바인딩','원본 코드 검토와 격리 테스트. 실행 파일 노출·타 객체 접근·허용 범위 밖 처리 없음'],
 ['CORS와 CSRF','API별 허용 origin·credentials, 상태 변경의 출처 검증','비허용 출처 접근과 위조 요청 차단, 정상 인증·제휴 흐름 유지'],
 ['내부 확산과 암호화','서비스 계정 DB·파일·네트워크 범위, 백업과 키 접근','최소 권한·키 분리·내부 구간 제한을 설정과 감사 기록으로 확인'],
 ['탐지와 복구','합성 이상 조회 이벤트, 경보 전달·세션 회수·복원 연습','합의한 탐지·대응·복구 목표 충족. 경보 ID·담당자 조치·복원 결과 기록']
],[3.0,6.5,7.7])
h('권한 검증 정답의 예',2)
table(['주체','A의 객체','B의 객체','관리 기능'],[
 ['비로그인','거부','거부','거부'],['사용자 A','허용 범위만','거부','거부'],['사용자 B','거부','허용 범위만','거부'],['제한된 운영자','업무 범위만','업무 범위만','별도 부여 범위만']
],[4.3,4.3,4.3,4.3])
p('거부는 상태코드만 확인하지 않는다. 응답의 데이터 노출, 실제 변경·내보내기 작업의 생성 여부와 우회 경로를 함께 확인한다. 운영에서 대량 조회나 성능 한계 시험을 수행하지 않는다.',size=9)

page();h('7 향후 진행 계획')
p('아래 기간은 착수와 자료 제공이 완료된 날을 D로 잡은 제안 일정이다. 실제 업무 범위·담당자·일정은 소유 부서가 확정해야 하며 약속된 완료일이 아니다.')
diagram('도식 3 자료 확보에서 운영 재검증까지의 진행 순서',[
 (15,20,215,110,'D부터 3영업일\n자산과 범위 확정\n소스와 테스트 환경','#E4EDF4'),
 (255,20,215,110,'D부터 10영업일\n브라우저 경로 수정\n인증 계약 확인','#E4EDF4'),
 (495,20,215,110,'D부터 20영업일\n권한과 반출 방어\n탐지 검증','#E4EDF4'),
 (735,20,215,110,'D부터 30영업일\n승인 후 운영 반영\n동일 조건 재검증','#E4EDF4')
],[(230,75,255,75,False),(470,75,495,75,False),(710,75,735,75,False)],85)
table(['단계와 권고 책임','진행할 사항','완료 산출물'],[
 ['1 범위와 증빙\nCISO·자산 소유자','전체 외부 접점, 데이터 흐름과 담당자 확정. 원본 소스·잠금 파일·설정·마스킹 로그와 A/B 계정 준비','대상 목록, 권한 범위, 운영과 테스트 차이, 자료 인수 목록'],
 ['2 우선 경로 개선\n프런트·인증·운영팀','F-07 텍스트 출력·F-10 목적지 제한, jQuery·Axios 정리. F-02·F-04·F-06 개선 및 F-08 계약 확인','원본에 통합한 패치, 영향 분석, 정상 흐름 회귀 결과, 배포·복구안'],
 ['3 핵심 통제 검증\n백엔드·DB·보안관제','서버 인증·객체 권한·청약과 파일·세션 회수·누적 반출 제한·최소 권한·탐지 검증','테스트별 기대·실제 결과, 비식별 증적, 미해결 위험과 책임자'],
 ['4 반영과 재검증\n변경 승인자·운영팀','변경 절차에 따라 적용. CDN 최종 응답과 실제 브라우저, 인증·제휴 흐름 재확인','적용 버전, 재검증 증적, 롤백 결과 또는 복구 가능성, 잔여 위험 승인'],
 ['5 정기 운영\n보안·개발 책임자','의존성·외부 접점 변경 시 재평가. 공식 위협정보와 운영 로그 대조 및 사고 훈련','자산 현행화, 예외 만료일, 재점검 일정, 경보·복구 훈련 기록']
],[3.8,7.8,5.6])
h('첫 착수 시 우선 확보할 자료',2)
p('프런트·백엔드 저장소와 빌드 버전, lock 파일과 SBOM, API 명세, 인증·토큰 계약, CDN·WAF·프록시·서버 설정, 외부 접점 및 서비스 계정 권한 목록, 마스킹한 로그, 테스트 환경과 전용 계정이 필요하다. 비밀키·실토큰·실고객 원문은 보고서 자료에 넣지 않는다.')

page();h('8 검증 현황과 종료 조건')
table(['검증 구분','기존 기록과 이번 확인','해석'],[
 ['원본 함수 로컬 재현','기존 HTML 함수 4건, URL 2건, 목적지 설정 5건 등 선택 구문 재현','동작 근거. 브라우저·서버 원격 공격 성공 아님'],
 ['수집 응답 조건 검사','기존 스냅샷 검사 7 통과·6 실패','프레임 제한·구버전 참조·내부 설정 관련 저장 응답 조건이 남아 있음'],
 ['수정 참고 코드 회귀','2026년 10월 29일 재실행 30 통과·0 실패\n상위 26건과 하위 4건 합계','로컬 예시 검증. 운영 서비스 수정 증거 아님'],
 ['실제 서비스 적용','원본 소스 통합·운영 배포·수정 후 재수집 미수행','완료 처리 불가'],
 ['실제 서버 방어','E5 절차만 수립. 인증·권한·반출·탐지 검사 미수행','안전성 판단을 위한 핵심 잔여 작업']
],[4.0,7.2,6.0])
h('개별 발견 사항의 종료 기준',2)
p('개발 완료와 보안 검증 완료를 분리한다. 항목별로 담당자, 수정 커밋과 배포 버전, 기대 결과, 실제 테스트 증적, 부작용 검토 및 승인자를 기록한다. 참고 코드가 존재하거나 테스트가 통과한 것만으로 F 항목을 종료하지 않는다.')
p('F-01·F-09는 실제 빌드와 실행 버전, F-02는 최종 응답 정책과 iframe 호환성, F-03·F-07은 입력부터 DOM까지의 안전한 출력, F-04는 공개 설정 제거, F-06은 비허용 데이터 저장·로그 제거, F-05·F-08은 실제 서버 수명주기, F-10은 실제 호출 경로와 전송 목적지 통제까지 확인한다.')
h('안전성 검토 회의의 의사결정',2)
p('최우선 경로를 수정하고 서버 인증·객체 권한·민감정보 접근·세션 회수·탐지를 검증한 뒤 잔여 위험을 재평가한다. 미해결 항목은 영향 범위, 보완 통제, 책임자, 기한과 재검증 방법을 명시해 승인 여부를 결정한다. 본 보고서 자체는 출시 승인서가 아니다.')
h('운영 관점의 안정성 확인',2)
p('보안 강화가 정상 업무를 막지 않도록 로그인·로그아웃·본인확인·전자서명·청약·제휴 iframe·캠페인 집계의 회귀 범위를 정한다. CSP, 라이브러리 교체, 쿠키 인증 전환은 연계 영향이 있어 단계 적용과 롤백 계획이 필요하다. 응답시간·오류율·복구목표는 운영 기준을 확보한 뒤 별도로 판정한다.')
p('현재는 가용성이나 성능 안정성을 계측하지 않았다. HTTPS 정상 관측과 정상 페이지 응답으로 업무 연속성·장애 복구까지 통과 처리하지 않는다.')

page();h('9 근거 자료와 참고 출처')
h('점검 산출물',2)
p('아래 경로는 저장소 루트 기준이다. F 항목은 문서와 해당 JSON·본문 증적을 함께 추적할 수 있다. 수집 SHA-256은 오편집 확인 수단이며 독립적인 서명이나 타임스탬프를 의미하지 않는다.',size=9)
table(['번호','점검 자료와 증적 경로'],[
 ['E1','audit/report.ko.md\naudit/evidence/responses.json\naudit/evidence/local-reproduction.json\naudit/evidence/snapshot-checks.json'],
 ['E2','audit/report-additional.ko.md\naudit/additional-evidence/responses.json\naudit/additional-evidence/local-verification.json'],
 ['E3','audit/intrusion-review.ko.md\naudit/path-evidence/responses.json\naudit/path-evidence/common/responses.json\naudit/path-evidence/local-path-verification.json'],
 ['E4','audit/financial-incidents-comparison.ko.md\naudit/financial-review-evidence/api-destination.json'],
 ['E5','audit/defense-validation.ko.md\naudit/README.md\naudit/fixes/ 및 audit/tests/']
],[1.5,15.7])
p('저장소  https://github.com/leeyonsei78/codex_pen_test\n증적 기준 커밋  a7d99d3d751e484267f25ffd5cfd03e48526c650',size=9)
h('외부 출처',2)
p('확인 기준일은 2026년 10월 29일이다. 사건 날짜와 발표·보도 날짜를 구분했으며, 조사 중인 항목은 후속 공지로 갱신해야 한다. 상세 피해는 아래 발표·보도 당시의 상태다.',size=9)
sources=[
 ('S1','금융위원회 2026년 10월 14일 금융권 침해사고 재발방지를 위한 제2차 긴급 상황대응 회의','https://www.fsc.go.kr/po010105/87927'),
 ('S2','CBS노컷뉴스 2026년 10월 2일 신한은행 대출모집인 조회 시스템과 개인정보 유출 보도','https://www.nocutnews.co.kr/news/6586151'),
 ('S3','뉴스1 2026년 10월 28일 미래에셋생명 홈페이지 차단 및 고객정보 유출 조사 보도','https://www.news1.kr/finance/insurance-card/6334573'),
 ('S4','조선비즈 2026년 10월 28일 대신증권 고객 38490명 정보 유출 발표 보도','https://biz.chosun.com/stock/stock_general/2026/10/28/YC2P4IBA55FHVE7IZFRDSBDXBE/'),
 ('S5','중앙일보 2026년 10월 28일 미래에셋자산운용 및 대신증권 등 침해 현황 보도','https://www.joongang.co.kr/article/25473563'),
]
for sid,title,url in sources:p(f'[{sid}] {title}\n{url}',size=8)
page();h('기술 기준과 보고서 활용')
for sid,title,url in [
 ('S6','jQuery 공식 보안 권고 CVE 2020 11022 및 11023','https://github.com/jquery/jquery/security/advisories/GHSA-gxr4-xjj5-5px2\nhttps://github.com/jquery/jquery/security/advisories/GHSA-jpcq-cgw6-v4j6'),
 ('S7','Axios 공식 0 28 0 릴리스와 교차 출처 XSRF 수정 PR 6028','https://github.com/axios/axios/releases/tag/v0.28.0\nhttps://github.com/axios/axios/pull/6028'),
 ('S8','OWASP Session Management 및 OAuth 2 0 보안 모범 관행 RFC 9700','https://cheatsheetseries.owasp.org/cheatsheets/Session_Management_Cheat_Sheet.html\nhttps://www.rfc-editor.org/rfc/rfc9700.html'),
 ('S9','OWASP DOM XSS 방어 및 인가 지침','https://cheatsheetseries.owasp.org/cheatsheets/DOM_based_XSS_Prevention_Cheat_Sheet.html\nhttps://cheatsheetseries.owasp.org/cheatsheets/Authorization_Cheat_Sheet.html')
]:p(f'[{sid}] {title}\n{url}',size=9)
h('용어와 판정 읽는 방법',2)
table(['표현','의미'],[
 ['확인','해당 수집 응답·코드 구문 또는 로컬 동작에서 근거를 확보함'],
 ['조건부 위험','추가 조건이 성립하면 피해로 이어질 수 있으나 전체 악용 경로는 미입증'],
 ['미확인 또는 미수행','자료·환경 또는 검사 범위의 한계로 판단하지 못함. 안전이나 결함으로 단정하지 않음'],
 ['참고 코드','통합 전 예시. 실제 서비스에 적용한 수정 코드와 구분'],
 ['origin','프로토콜과 호스트 및 포트로 구분되는 출처. URL에 특정 단어가 있는지와 다름'],
 ['객체 권한','요청자가 해당 고객·문서·계약 등 개별 데이터에 접근할 권한이 있는지 서버에서 검사']
],[4.2,13])
h('문서와 도식 갱신',2)
p('도식 1부터 3은 편집 가능한 Word 도형이며 비교표와 권한표도 Word 표로 작성했다. 도형 문구를 수정할 때 확인과 미확인 경계를 함께 갱신한다. 새 증적이 확보되면 해당 F 번호와 E 출처, 적용 버전 및 재검증 날짜를 연결해 판정을 변경한다.')
p('권고 통제의 기술적 배경은 S6부터 S9 및 기존 점검 문서의 공식 지침을 따른다. 법령 준수 여부나 제재 책임은 본 기술 점검만으로 판정하지 않는다.')

out=OUT/'라이나생명_보안점검_통합보고서.docx'
doc.save(out)
print(out)
print('tables',len(doc.tables),'diagrams',3)
